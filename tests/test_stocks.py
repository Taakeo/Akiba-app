def test_stocks_requires_permission(client, login_seller):
    assert client.get("/stocks/").status_code == 403
    assert client.get("/stocks/ajustement").status_code == 403
    assert client.get("/stocks/inventaires").status_code == 403


def test_stocks_index_lists_produit(client, login_admin, catalogue):
    response = client.get("/stocks/")
    assert response.status_code == 200
    assert b"Tablette Chocolat 70%" in response.data


def test_stocks_index_expose_les_attributs_de_filtre(client, login_admin, catalogue):
    # Les filtres live (JS) s'appuient sur ces data-* : verrouille leur
    # présence pour que le JS ne se retrouve jamais silencieusement inerte.
    response = client.get("/stocks/")
    assert response.status_code == 200
    html = response.data.decode()
    assert 'data-statut-stock="ok"' in html
    assert f'data-poste-id="{catalogue["poste_id"]}"' in html
    assert f'data-categorie-id="{catalogue["categorie_id"]}"' in html
    assert 'id="filtre-recherche"' in html


def test_ajustement_sortie_perte_decrements_stock(client, login_admin, catalogue):
    response = client.post(
        "/stocks/ajustement",
        data={
            "produit_id": str(catalogue["produit_id"]),
            "type_mouvement": "sortie",
            "motif_sortie": "perte",
            "motif_entree": "don",
            "quantite": "3",
            "commentaire": "Casse",
        },
    )
    assert response.status_code == 302

    from app.extensions import db
    from app.models import MouvementStock, Produit

    produit = db.session.get(Produit, catalogue["produit_id"])
    assert produit.stock_quantite == 7  # 10 - 3

    mouvement = MouvementStock.query.filter_by(produit_id=produit.id).first()
    assert mouvement.motif == "perte"
    assert mouvement.type_mouvement == "sortie"


def test_ajustement_sortie_insuffisante_bloquee(client, login_admin, catalogue):
    response = client.post(
        "/stocks/ajustement",
        data={
            "produit_id": str(catalogue["produit_id"]),
            "type_mouvement": "sortie",
            "motif_sortie": "perte",
            "motif_entree": "don",
            "quantite": "99",
        },
    )
    assert response.status_code == 200

    from app.extensions import db
    from app.models import Produit

    produit = db.session.get(Produit, catalogue["produit_id"])
    assert produit.stock_quantite == 10  # inchangé


def test_inventaire_general_cloture_applique_ecart(client, login_admin, catalogue):
    response = client.post("/stocks/inventaires/nouveau", data={"type_inventaire": "general"})
    assert response.status_code == 302
    inventaire_url = response.headers["Location"]

    from app.models import Inventaire, InventaireLigne

    inventaire = Inventaire.query.order_by(Inventaire.id.desc()).first()
    ligne = InventaireLigne.query.filter_by(inventaire_id=inventaire.id, produit_id=catalogue["produit_id"]).first()
    assert ligne.stock_theorique == 10

    response = client.post(inventaire_url, data={f"reel_{ligne.id}": "8"})
    assert response.status_code == 302

    from app.extensions import db
    from app.models import MouvementStock, Produit

    produit = db.session.get(Produit, catalogue["produit_id"])
    assert produit.stock_quantite == 8

    mouvement = MouvementStock.query.filter_by(produit_id=produit.id, motif="correction").first()
    assert mouvement is not None
    assert mouvement.type_mouvement == "sortie"
    assert mouvement.quantite == 2

    db.session.refresh(inventaire)
    assert inventaire.statut == "cloture"


# --- Export/réimport Excel de l'inventaire ------------------------------------


def test_export_inventaire_route_requiert_permission(client, login_seller):
    assert client.get("/stocks/export.xlsx").status_code == 403


def test_export_inventaire_liste_le_produit_du_catalogue(client, login_admin, catalogue):
    response = client.get("/stocks/export.xlsx")
    assert response.status_code == 200

    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.data))
    ws = wb.active
    lignes = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(lignes) == 1
    assert lignes[0][0] == catalogue["produit_id"]
    assert lignes[0][1] == "Tablette Chocolat 70%"
    assert lignes[0][5] == 10  # stock actuel


def _workbook_inventaire_bytes(rows):
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "Nom du produit", "Poste", "Catégorie", "Unité", "Stock actuel", "Seuil d'alerte", "Code-barres"])
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def test_importer_inventaire_corrige_le_stock_et_trace_le_mouvement(app, db, catalogue):
    from app.admin.import_excel import importer_inventaire_excel
    from app.models import MouvementStock, Produit, SubProfile

    with app.app_context():
        from app.models import Profile

        profile = Profile(code="administrateur", name="Administrateur", icon="shield_person")
        profile.permissions = ["*"]
        db.session.add(profile)
        db.session.commit()
        current_user = SubProfile(profile_id=profile.id, full_name="Admin")
        current_user.set_pin("9999")
        db.session.add(current_user)
        db.session.commit()

        fichier = _workbook_inventaire_bytes(
            [[catalogue["produit_id"], "Tablette Chocolat 70%", "Boutique", "Boutique", "unité", 7, 2, ""]]
        )
        resultat = importer_inventaire_excel(fichier, current_user)

        assert len(resultat["maj"]) == 1
        assert resultat["erreurs"] == []

        produit = db.session.get(Produit, catalogue["produit_id"])
        assert produit.stock_quantite == 7  # 10 -> 7

        mouvement = MouvementStock.query.filter_by(produit_id=produit.id, motif="correction").first()
        assert mouvement is not None
        assert mouvement.type_mouvement == "sortie"
        assert mouvement.quantite == 3


def test_importer_inventaire_id_inconnu_est_une_erreur(app, db, catalogue):
    from app.admin.import_excel import importer_inventaire_excel
    from app.models import Profile, SubProfile

    with app.app_context():
        profile = Profile(code="administrateur", name="Administrateur", icon="shield_person")
        profile.permissions = ["*"]
        db.session.add(profile)
        db.session.commit()
        current_user = SubProfile(profile_id=profile.id, full_name="Admin")
        current_user.set_pin("9999")
        db.session.add(current_user)
        db.session.commit()

        fichier = _workbook_inventaire_bytes([[999999, "Inconnu", "Boutique", "Boutique", "unité", 5, "", ""]])
        resultat = importer_inventaire_excel(fichier, current_user)

        assert resultat["maj"] == []
        assert len(resultat["erreurs"]) == 1
