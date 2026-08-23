from datetime import date, datetime, timezone


def _creer_vente(db, catalogue, quantite=2, prix_unitaire=8000):
    from app.models import CaisseSession, LigneVente, Vente

    session = CaisseSession(
        compte_financier_id=catalogue["caisse_id"], fond_ouverture=0, ouverte_par_nom="Admin"
    )
    db.session.add(session)
    db.session.flush()

    vente = Vente(
        caisse_session_id=session.id,
        type_tarif_id=catalogue["type_tarif_id"],
        created_by_name="Admin",
        sous_total=quantite * prix_unitaire,
        total=quantite * prix_unitaire,
    )
    db.session.add(vente)
    db.session.flush()

    from app.models import Produit

    produit = db.session.get(Produit, catalogue["produit_id"])
    ligne = LigneVente(
        vente_id=vente.id,
        produit_id=produit.id,
        produit_nom=produit.name,
        poste_id=produit.poste_id,
        categorie_id=produit.categorie_id,
        quantite=quantite,
        prix_unitaire=prix_unitaire,
        total_ligne=quantite * prix_unitaire,
    )
    db.session.add(ligne)
    db.session.commit()
    return vente


def test_rapport_ventes_regroupe_par_produit(client, login_admin, catalogue, app, db):
    with app.app_context():
        _creer_vente(db, catalogue, quantite=3, prix_unitaire=8000)

    today = date.today().isoformat()
    response = client.get(f"/rapports/ventes?debut={today}&fin={today}&group_by=produit")
    assert response.status_code == 200
    assert b"Tablette Chocolat 70%" in response.data
    assert "24 000".encode() in response.data or b"24000" in response.data


def test_rapport_ventes_export_xlsx(client, login_admin, catalogue, app, db):
    with app.app_context():
        _creer_vente(db, catalogue, quantite=1, prix_unitaire=8000)

    today = date.today().isoformat()
    response = client.get(f"/rapports/ventes/export.xlsx?debut={today}&fin={today}")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert len(response.data) > 0


def test_rapport_ventes_export_pdf(client, login_admin, catalogue, app, db):
    with app.app_context():
        _creer_vente(db, catalogue, quantite=1, prix_unitaire=8000)

    today = date.today().isoformat()
    response = client.get(f"/rapports/ventes/export.pdf?debut={today}&fin={today}")
    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert response.data[:4] == b"%PDF"


def test_rapports_requires_permission(client, login_seller):
    response = client.get("/rapports/")
    assert response.status_code == 403


def test_dashboard_shows_ca_jour(client, login_admin, catalogue, app, db):
    with app.app_context():
        _creer_vente(db, catalogue, quantite=2, prix_unitaire=8000)

    response = client.get("/")
    assert response.status_code == 200
    assert "16 000".encode() in response.data


def test_rapport_production_regroupe_par_produit(client, login_admin, catalogue, app, db):
    from datetime import date as date_cls

    with app.app_context():
        from app.models import Fabrication, Produit

        produit = db.session.get(Produit, catalogue["produit_id"])
        db.session.add(
            Fabrication(
                produit_id=produit.id,
                quantite=25,
                responsable_nom="Lala",
                date_fabrication=date_cls.today(),
                created_by_name="Admin",
            )
        )
        db.session.commit()

    today = date.today().isoformat()
    response = client.get(f"/rapports/production?debut={today}&fin={today}")
    assert response.status_code == 200
    assert b"Tablette Chocolat 70%" in response.data
    assert b"25" in response.data


def test_rapport_rh_soustrait_les_retenues(client, login_admin, catalogue, app, db):
    from datetime import date as date_cls

    with app.app_context():
        from app.models import RemunerationSalarie, Salarie

        salarie = Salarie(nom="Lala Rakoto", poste_id=catalogue["poste_id"])
        db.session.add(salarie)
        db.session.flush()
        db.session.add_all(
            [
                RemunerationSalarie(
                    salarie_id=salarie.id,
                    type_remuneration="salaire_mensuel",
                    montant=300000,
                    date_versement=date_cls.today(),
                    created_by_name="Admin",
                ),
                RemunerationSalarie(
                    salarie_id=salarie.id,
                    type_remuneration="retenue",
                    montant=5000,
                    date_versement=date_cls.today(),
                    created_by_name="Admin",
                ),
            ]
        )
        db.session.commit()

    today = date.today().isoformat()
    response = client.get(f"/rapports/rh?debut={today}&fin={today}&group_by=salarie")
    assert response.status_code == 200
    assert b"Lala Rakoto" in response.data
    assert "295 000".encode() in response.data  # 300000 - 5000


def test_rapports_production_rh_requires_permission(client, login_seller):
    assert client.get("/rapports/production").status_code == 403
    assert client.get("/rapports/rh").status_code == 403


def test_rapport_ventes_groupe_par_poste_categorie_projet(client, login_admin, catalogue, app, db):
    # Bug latent corrigé au passage (LigneVente.poste/categorie/projet
    # manquaient) : ces regroupements n'étaient jamais exercés par les tests
    # existants (seuls "produit"/"date" l'étaient) et levaient une
    # AttributeError dès qu'on les utilisait réellement.
    with app.app_context():
        _creer_vente(db, catalogue, quantite=1, prix_unitaire=8000)

    today = date.today().isoformat()
    for group_by in ("poste", "categorie", "projet"):
        response = client.get(f"/rapports/ventes?debut={today}&fin={today}&group_by={group_by}")
        assert response.status_code == 200


def test_rapport_ventes_groupe_par_date_ordre_chronologique(client, login_admin, catalogue, app, db):
    with app.app_context():
        _creer_vente(db, catalogue, quantite=1, prix_unitaire=8000)

    today = date.today().isoformat()
    response = client.get(f"/rapports/ventes?debut={today}&fin={today}&group_by=date")
    assert response.status_code == 200
    assert today.encode() in response.data


def test_dashboard_affiche_top_produits(client, login_admin, catalogue, app, db):
    with app.app_context():
        _creer_vente(db, catalogue, quantite=4, prix_unitaire=8000)

    response = client.get("/")
    assert response.status_code == 200
    assert b"Produits les plus vendus" in response.data
    assert b"Tablette Chocolat 70%" in response.data
    assert b"4 vendu" in response.data


def _preparer_rapport_total(app, db, catalogue):
    """Une ligne de chaque source (vente PDV, vente externe, achat, RH payée
    et RH retenue) sur la période du jour — pour vérifier que le rapport
    total les combine toutes correctement."""
    from datetime import date as date_cls

    from app.models import (
        Achat,
        CompteFinancier,
        RemunerationSalarie,
        Salarie,
        VenteExterne,
    )

    with app.app_context():
        _creer_vente(db, catalogue, quantite=2, prix_unitaire=8000)  # recette 16 000

        compte_akiba = CompteFinancier(name="Compte Akiba", devise="Ar", is_compte_akiba=True)
        db.session.add(compte_akiba)
        db.session.flush()

        db.session.add(
            VenteExterne(
                client_nom="Jean Rakoto",
                date_vente=date_cls.today(),
                poste_id=catalogue["poste_id"],
                categorie_id=catalogue["categorie_id"],
                montant_total=15000,  # recette
                moyen_paiement_id=catalogue["moyen_paiement_id"],
                created_by_name="Admin",
            )
        )

        db.session.add(
            Achat(
                type_achat="depense",
                date_achat=date_cls.today(),
                poste_id=catalogue["poste_id"],
                categorie_id=catalogue["categorie_id"],
                montant_total=5000,  # dépense
                moyen_paiement_id=catalogue["moyen_paiement_id"],
                created_by_name="Admin",
            )
        )

        salarie = Salarie(nom="Lala Rakoto", poste_id=catalogue["poste_id"])
        db.session.add(salarie)
        db.session.flush()
        db.session.add_all(
            [
                RemunerationSalarie(
                    salarie_id=salarie.id,
                    type_remuneration="salaire_mensuel",
                    montant=300000,  # dépense (payée, moyen renseigné)
                    date_versement=date_cls.today(),
                    moyen_paiement_id=catalogue["moyen_paiement_id"],
                    created_by_name="Admin",
                ),
                RemunerationSalarie(
                    salarie_id=salarie.id,
                    type_remuneration="retenue",
                    montant=5000,  # jamais un effet sur les comptes -> exclue
                    date_versement=date_cls.today(),
                    created_by_name="Admin",
                ),
            ]
        )
        db.session.commit()


def test_rapport_total_combine_les_quatre_sources(client, login_admin, catalogue, app, db):
    _preparer_rapport_total(app, db, catalogue)

    today = date.today().isoformat()
    response = client.get(f"/rapports/total?debut={today}&fin={today}")
    assert response.status_code == 200
    # Recettes : 16 000 (vente PDV) + 15 000 (vente externe) = 31 000
    assert "31 000".encode() in response.data
    # Dépenses : 5 000 (achat) + 300 000 (salaire payé) = 305 000 — la retenue
    # (5 000, jamais versée) n'a aucun effet et ne doit pas apparaître ici.
    assert "305 000".encode() in response.data
    # Solde : 31 000 - 305 000 = -274 000
    assert "-274 000".encode() in response.data


def test_rapport_total_export_xlsx_trois_feuilles(client, login_admin, catalogue, app, db):
    _preparer_rapport_total(app, db, catalogue)

    today = date.today().isoformat()
    response = client.get(f"/rapports/total/export.xlsx?debut={today}&fin={today}")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames == ["Saisie", "Listes", "BILAN"]
    # "Listes" est une feuille technique, jamais destinée à être modifiée à la
    # main — cachée pour ne pas embrouiller l'utilisateur.
    assert wb["Listes"].sheet_state == "hidden"


def test_rapport_total_export_xlsx_listes_deroulantes_en_cascade(client, login_admin, catalogue, app, db):
    # Vérifie le vrai mécanisme demandé par l'utilisateur (listes déroulantes
    # Poste -> Catégorie -> Sous-catégorie comme dans le fichier Excel de
    # référence), pas seulement que le fichier s'ouvre.
    _preparer_rapport_total(app, db, catalogue)

    today = date.today().isoformat()
    response = client.get(f"/rapports/total/export.xlsx?debut={today}&fin={today}")

    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.data))

    # Plage nommée listant tous les postes, alimentant le premier niveau de la
    # cascade (colonne Poste de Saisie).
    assert "ListePostes" in wb.defined_names.keys()

    # Une plage nommée par poste, listant ses catégories (ex. "Boutique" du
    # catalogue de test -> plage nommée "Cat_Boutique").
    noms_plages = set(wb.defined_names.keys())
    assert any(nom.startswith("Cat_") for nom in noms_plages)

    ws_saisie = wb["Saisie"]
    # Colonne B (Poste) : liste simple sourcée sur la plage nommée globale.
    dv_poste = next(dv for dv in ws_saisie.data_validations.dataValidation if str(dv.sqref).startswith("B6"))
    assert dv_poste.formula1 == "ListePostes"
    # Colonne C (Catégorie) : dépend de la colonne technique cachée O (elle-
    # même dérivée du Poste choisi en colonne B) — cascade en INDIRECT().
    dv_categorie = next(dv for dv in ws_saisie.data_validations.dataValidation if str(dv.sqref).startswith("C6"))
    assert dv_categorie.formula1 == "INDIRECT(O6)"

    # Colonnes techniques N à Q cachées (support de la cascade, pas destinées
    # à être vues/modifiées).
    for lettre in ("N", "O", "P", "Q"):
        assert ws_saisie.column_dimensions[lettre].hidden is True


def test_rapport_total_filtre_par_poste(client, login_admin, catalogue, app, db):
    from app.models import Poste

    with app.app_context():
        autre_poste = Poste(name="Agriculture")
        db.session.add(autre_poste)
        db.session.commit()
        autre_poste_id = autre_poste.id

    _preparer_rapport_total(app, db, catalogue)

    today = date.today().isoformat()
    response = client.get(f"/rapports/total?debut={today}&fin={today}&poste_id={autre_poste_id}")
    assert response.status_code == 200
    assert b"Jean Rakoto" not in response.data  # filtré : appartient au poste du catalogue, pas "Agriculture"


def test_rapports_total_requires_permission(client, login_seller):
    assert client.get("/rapports/total").status_code == 403
